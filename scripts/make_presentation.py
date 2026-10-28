"""Generează prezentarea de pitch: prezentare_ORB.pptx (16:9, culorile din frontend).

Structura urmează README §13 (pitch ~4:30); cifrele vin din notes/stefan.md și data/benchmark/.
Fiecare slide are în „Note” textul de spus. Ce e marcat [DE COMPLETAT] vine de la mentori / echipă.

Rulare (din rădăcina repo-ului):
  pip install python-pptx
  python scripts/make_presentation.py [--lang ro|en] [--out prezentare_ORB.pptx]
  --lang en: pitch-ul scurt în engleză, 6 slide-uri după notes/plan.md și notes/pitch.md → prezentare_ORB_EN.pptx
"""
import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]

# aceleași culori ca în frontend/src/index.css
INK = RGBColor(0x1F, 0x23, 0x28)
INK2 = RGBColor(0x4E, 0x57, 0x64)
INK3 = RGBColor(0x6B, 0x74, 0x80)
BRAND = RGBColor(0xFF, 0xE6, 0x00)
CANVAS = RGBColor(0xED, 0xEF, 0xF2)
LINE = RGBColor(0xD5, 0xD9, 0xDF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DANGER = RGBColor(0xB4, 0x23, 0x18)
DANGER_BG = RGBColor(0xFE, 0xF3, 0xF2)
OK = RGBColor(0x06, 0x76, 0x47)
OK_BG = RGBColor(0xEC, 0xFD, 0xF3)
WARN = RGBColor(0xB5, 0x47, 0x08)
WARN_BG = RGBColor(0xFF, 0xFA, 0xEB)
NIGHT = RGBColor(0x16, 0x19, 0x1E)

FONT = "Arial"  # există și în Google Slides
W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.7)  # marginea laterală


# ── unelte ────────────────────────────────────────────────────────────────────

def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, content, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """`content`: un șir sau o listă de paragrafe; un paragraf poate fi (text, {size, color, bold})."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, item in enumerate(content if isinstance(content, list) else [content]):
        t, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(opt.get("after", 6))
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = FONT
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", bold)
        f.color.rgb = opt.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=22, color=INK):
    return text(slide, x, y, w, h, [(f"•  {t}", {"after": 12}) for t in items], size=size, color=color)


def base(prs, title, kicker=None, notes="", dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # gol
    box(s, 0, 0, W, H, fill=NIGHT if dark else CANVAS)
    box(s, 0, 0, W, Inches(0.09), fill=BRAND)
    if kicker:
        text(s, MX, Inches(0.4), Inches(10), Inches(0.4), kicker.upper(), size=14, bold=True,
             color=BRAND if dark else INK3)
    if title:
        text(s, MX, Inches(0.75), W - 2 * MX, Inches(1), title, size=32, bold=True, color=WHITE if dark else INK)
    # subsolul: marca ORB + numărul slide-ului
    n = len(prs.slides)
    text(s, MX, H - Inches(0.5), Inches(4), Inches(0.3), "ORB · Raiffeisen Bank: Operational Efficiency", size=11,
         color=INK3)
    text(s, W - MX - Inches(1), H - Inches(0.5), Inches(1), Inches(0.3), str(n), size=11, color=INK3,
         align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def card(slide, x, y, w, h, title, body, accent=INK, fill=WHITE, title_size=20, body_size=16):
    box(slide, x, y, w, h, fill=fill, line=LINE)
    box(slide, x, y, Inches(0.08), h, fill=accent)
    text(slide, x + Inches(0.25), y + Inches(0.15), w - Inches(0.4), Inches(0.5), title, size=title_size, bold=True)
    text(slide, x + Inches(0.25), y + Inches(0.65), w - Inches(0.4), h - Inches(0.75), body, size=body_size,
         color=INK2)


def big_number(slide, x, y, w, value, label, color=INK):
    text(slide, x, y, w, Inches(1.0), value, size=40, bold=True, color=color)
    text(slide, x, y + Inches(0.95), w, Inches(1.2), label, size=16, color=INK2)


def table(slide, x, y, w, rows, col_w, header=True, size=16, row_h=Inches(0.5)):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows))
    t = shape.table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Emu(int(w * cw))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = INK if header and i == 0 else WHITE if i % 2 else CANVAS
            cell.margin_left = cell.margin_right = Inches(0.12)
            p = cell.text_frame.paragraphs[0]
            p.text = ""
            r = p.add_run()
            r.text = val
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = header and i == 0
            r.font.color.rgb = WHITE if header and i == 0 else INK
    return t


# ── slide-uri ─────────────────────────────────────────────────────────────────

def s_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, W, H, fill=NIGHT)
    box(s, 0, 0, W, Inches(0.12), fill=BRAND)
    box(s, MX, Inches(1.6), Inches(1.2), Inches(1.2), fill=BRAND)
    text(s, MX, Inches(1.6), Inches(1.2), Inches(1.2), "O", size=60, bold=True, color=NIGHT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX + Inches(1.5), Inches(1.6), Inches(8), Inches(1.2), "ORB", size=72, bold=True, color=WHITE,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, MX, Inches(3.2), Inches(11.5), Inches(1.6),
         "Agenți AI care verifică documentele băncii fără să le citească", size=40, bold=True, color=WHITE)
    text(s, MX, Inches(4.9), Inches(11), Inches(0.6),
         "Hackathon · Raiffeisen Bank: Operational Efficiency", size=22, color=BRAND)
    text(s, MX, Inches(5.6), Inches(11), Inches(0.6),
         "Samuel · Narcis · Stefan · Erin · Lorena", size=20, color=RGBColor(0xA8, 0xB0, 0xBC))
    # locul pentru QR-ul Arenei (README §13: apare de la început)
    box(s, W - MX - Inches(2.4), Inches(4.4), Inches(2.4), Inches(2.4), fill=WHITE)
    text(s, W - MX - Inches(2.4), Inches(4.4), Inches(2.4), Inches(2.4),
         ["QR Arena", ("[DE COMPLETAT: lipește QR-ul]", {"size": 11, "color": INK3})], size=18, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    s.notes_slide.notes_text_frame.text = (
        "Scanați QR-ul: puteți încerca să păcăliți agenții noștri în Arena cât timp vorbim.")


def s_maria(prs):
    s = base(prs, "Maria Ionescu, 67 de ani, pensionară din Ploiești", "Problema",
             notes="Un escroc găsește o poză cu buletinul Mariei și cere online un credit de 45.000 lei pe numele ei. "
                   "Toate cele trei documente sunt false — și arată credibil. Azi, un om din back-office le verifică "
                   "de mână, unul câte unul.")
    text(s, MX, Inches(1.8), Inches(6.2), Inches(1.2),
         "Un escroc cere online un credit de 45.000 lei pe numele ei, cu:", size=24, color=INK2)
    items = [("Buletinul Mariei", "cu poza înlocuită de un chip generat cu AI"),
             ("Fluturaș de salariu", "făcut cu un LLM: „contabilă, 8.500 lei net”"),
             ("Extras de cont", "editat: soldul final umflat")]
    for i, (t, b) in enumerate(items):
        card(s, MX, Inches(2.9) + i * Inches(1.25), Inches(6.2), Inches(1.1), t, b, accent=DANGER)
    box(s, Inches(7.4), Inches(1.8), Inches(5.2), Inches(4.6), fill=WHITE, line=LINE)
    text(s, Inches(7.7), Inches(2.1), Inches(4.6), Inches(4),
         [("Azi", {"size": 16, "bold": True, "color": INK3}),
          ("Verificarea e manuală, document cu document.", {"size": 24, "bold": True}),
          ("Falsurile făcute cu AI arată tot mai bine, iar fraudele în rețea nu se văd dintr-un singur dosar.",
           {"size": 18, "color": INK2}),
          ("Și datele clienților nu au voie să iasă din bancă.", {"size": 18, "color": INK2})], size=18)


def s_solution(prs):
    s = base(prs, "Un birou virtual de agenți AI pentru back-office", "Soluția",
             notes="ORB citește documentele dintr-o cerere de credit, le verifică, prinde falsurile și leagă dosarele "
                   "între ele. Nu aprobă credite: pregătește dosarul cu dovezi, iar omul decide.")
    cols = [("Citește", "buletin, fluturaș, extras, adeverință — fiecare valoare legată de locul ei din document"),
            ("Verifică", "taxele, soldurile, firma la ANAF, CNP vs. buletin vs. MRZ, vârsta, IBAN-ul"),
            ("Prinde falsuri", "editări (ELA, structură PDF, metadate) și rețele: același telefon pe 5 nume"),
            ("Omul decide", "verdict cu scor de risc, chenar pe dovadă și jurnal de audit")]
    cw = (W - 2 * MX - Inches(0.6)) / 4
    for i, (t, b) in enumerate(cols):
        card(s, MX + i * (cw + Inches(0.2)), Inches(2.0), cw, Inches(2.6), t, b,
             accent=OK if i == 3 else INK, body_size=17)
    box(s, MX, Inches(5.0), W - 2 * MX, Inches(1.4), fill=WARN_BG, line=WARN)
    text(s, MX + Inches(0.3), Inches(5.1), W - 2 * MX - Inches(0.6), Inches(1.2),
         "Fără ca datele clienților să iasă din bancă: AI-ul din cloud vede doar [PERSOANA_1], [CNP_1] — "
         "iar un buton trece totul pe modul 100% local.", size=20, bold=True, color=WARN,
         anchor=MSO_ANCHOR.MIDDLE)


def s_agents(prs):
    s = base(prs, "Șase agenți, în ~15 secunde pe dosar", "Cum funcționează",
             notes="Paznicul ascunde datele personale local, înainte de orice AI. Funcționarul extrage datele, "
                   "Verificatorul reface calculele, Detectorul caută editări, Detectivul leagă dosarele, "
                   "iar Compliance dă verdictul din dovezi. Omul ia decizia.")
    agents = [("Paznic", "OCR, ascunde datele personale, seif local"),
              ("Funcționar", "extrage datele, fără voie să inventeze"),
              ("Verificator", "taxe, extras, ANAF, verificări RO"),
              ("Detector", "editări și documente făcute cu AI"),
              ("Detectiv", "rețele de fraudă între dosare"),
              ("Compliance", "scor din dovezi + verdict")]
    n = len(agents)
    gap = Inches(0.25)
    bw = (W - 2 * MX - gap * (n - 1)) / n
    y = Inches(2.3)
    for i, (t, b) in enumerate(agents):
        x = MX + i * (bw + gap)
        box(s, x, y, bw, Inches(2.4), fill=WHITE, line=LINE)
        box(s, x, y, bw, Inches(0.55), fill=INK)
        text(s, x, y, bw, Inches(0.55), f"{i + 1}  {t}", size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + Inches(0.12), y + Inches(0.7), bw - Inches(0.24), Inches(1.6), b, size=16, color=INK2,
             align=PP_ALIGN.CENTER)
    box(s, MX, Inches(5.1), W - 2 * MX, Inches(1.1), fill=OK_BG, line=OK)
    text(s, MX + Inches(0.3), Inches(5.1), W - 2 * MX - Inches(0.6), Inches(1.1),
         "👤  Angajatul băncii decide — cu chenarul pe dovadă, nu cu o cutie neagră.", size=22, bold=True,
         color=OK, anchor=MSO_ANCHOR.MIDDLE)


def s_case(prs):
    s = base(prs, "Dosarul Mariei: scor de risc 96 → investigație", "Demo live",
             notes="În demo arătăm trei momente: ce a văzut AI-ul lângă original, harta de editare care se aprinde "
                   "pe extras și graful care arată rețeaua. Apoi angajatul respinge cererea, iar banca o anunță "
                   "pe Maria pe telefonul din evidențele ei, nu pe cel din cerere.")
    rows = [["Severitate", "Ce a găsit ORB"],
            ["critic", "Angajatorul figurează inactiv la ANAF"],
            ["ridicat", "Netul calculat (8.307 lei) diferă de cel declarat (8.500 lei)"],
            ["ridicat", "Soldul final (26.480 lei) nu iese din tranzacții (2.851,50 lei)"],
            ["mediu", "Netul de pe fluturaș nu apare ca încasare în extras"],
            ["mediu", "Solicitantă de 67 de ani declarată angajată activ"],
            ["scăzut", "Extrasul arată o încasare de pensie"],
            ["rețea", "Telefonul din cerere apare în alte 5 cereri, pe 5 nume diferite"]]
    table(s, MX, Inches(1.9), Inches(8.3), rows, [0.22, 0.78], size=16)
    box(s, Inches(9.4), Inches(1.9), Inches(3.25), Inches(4.4), fill=DANGER_BG, line=DANGER)
    text(s, Inches(9.6), Inches(2.1), Inches(2.9), Inches(4),
         [("Verdict", {"size": 16, "bold": True, "color": DANGER}),
          ("96", {"size": 80, "bold": True, "color": DANGER, "after": 0}),
          ("risc ridicat", {"size": 18, "color": DANGER}),
          ("Furt de identitate, rețea suspectă → investigație", {"size": 18, "bold": True, "color": INK})])


def s_privacy(prs):
    s = base(prs, "Datele clienților nu ies din bancă", "Confidențialitate",
             notes="Spunem exact ce facem: identificatorii nu părăsesc serverul băncii; ce ajunge la LLM e "
                   "pseudonimizat, iar o poartă blochează orice payload în care apare o valoare din seif. Nu spunem "
                   "că AI-ul nu vede nimic — vede sume și date, pseudonimizate. Demo: butonul de mod local, cablul "
                   "scos, contorul „0 apeluri externe”.")
    items = [("Pseudonimizare locală", "CNP, IBAN, nume, MRZ, poza → [CNP_1], [PERSOANA_1]. Pozele buletinului "
                                       "nu pleacă niciodată.", INK),
             ("Poartă anti-PII", "orice payload către LLM cu o valoare din seif e blocat înainte să plece.", INK),
             ("Mod 100% local", "un buton: LLM pe serverul băncii (Ollama), zero apeluri externe.", OK),
             ("Criptare homomorfică", "un verificator terț reface calculul netului pe numere criptate, "
                                      "fără să vadă salariul (CKKS, ~0,3 s).", INK)]
    cw = (W - 2 * MX - Inches(0.3)) / 2
    for i, (t, b, a) in enumerate(items):
        card(s, MX + (i % 2) * (cw + Inches(0.3)), Inches(1.9) + (i // 2) * Inches(2.2), cw, Inches(1.95), t, b,
             accent=a, body_size=18)


def s_zk(prs):
    s = base(prs, "Clientul își dovedește venitul fără să-l arate", "Zero-knowledge",
             notes="Întrebarea băncii: venitul net acoperă rata? Angajatorul semnează un angajament, clientul face "
                   "o dovadă pe telefon, banca o verifică și răspunde DA sau NU — fără să afle salariul. Arătăm și "
                   "că un client care minte, coboară pragul, falsifică semnătura sau retrimite o dovadă primește NU.")
    steps = [("Banca", "cere dovada: venit ≥ rată / 0,4, cu nonce unic"),
             ("Angajatorul", "semnează angajamentul (ECDSA); salariul merge doar la client"),
             ("Clientul", "generează dovada ZK (Noir + Barretenberg)"),
             ("Banca", "verifică semnătura și dovada → DA / NU")]
    sw = (W - 2 * MX - Inches(0.9)) / 4
    for i, (t, b) in enumerate(steps):
        x = MX + i * (sw + Inches(0.3))
        card(s, x, Inches(1.9), sw, Inches(2.0), f"{i + 1}. {t}", b, accent=BRAND, body_size=16)
    box(s, MX, Inches(4.25), Inches(5.6), Inches(2.1), fill=OK_BG, line=OK)
    text(s, MX + Inches(0.3), Inches(4.4), Inches(5.1), Inches(1.9),
         [("Ce vede banca", {"size": 16, "bold": True, "color": OK}),
          ("CUI-ul angajatorului, luna, angajamentul și o dovadă de câțiva KB.", {"size": 18}),
          ("Salariul: nicăieri.", {"size": 22, "bold": True, "color": OK})])
    box(s, Inches(6.6), Inches(4.25), W - MX - Inches(6.6), Inches(2.1), fill=DANGER_BG, line=DANGER)
    text(s, Inches(6.9), Inches(4.4), W - MX - Inches(7.2), Inches(1.9),
         [("Atacuri respinse → NU", {"size": 16, "bold": True, "color": DANGER}),
          ("minte la salariu · coboară pragul · semnătură falsă · retrimite dovada", {"size": 20})])


def s_arena(prs):
    s = base(prs, "Arena: oricine poate păcăli agenții", "AI contra AI",
             notes="Juriul a jucat din telefon cât am vorbit. Fiecare fals care trece devine o regulă nouă: "
                   "propusă de LLM, testată pe dosarele vechi (backtest), aprobată de un om. [Aici: rezultatele "
                   "juriului de pe /arena/live și „acest fals a devenit regula #N”.]")
    flow = ["Faci un fals pe telefon (Forge Studio)", "Pipeline-ul real îl verifică",
            "Dacă trece: LLM-ul propune o regulă",
            "Backtest pe dosarele vechi", "Un om aprobă → regula intră în Verificator"]
    for i, t in enumerate(flow):
        y = Inches(1.9) + i * Inches(0.85)
        box(s, MX, y, Inches(0.6), Inches(0.6), fill=BRAND, shape=MSO_SHAPE.OVAL)
        text(s, MX, y, Inches(0.6), Inches(0.6), str(i + 1), size=20, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, MX + Inches(0.85), y, Inches(6.9), Inches(0.6), t, size=19, anchor=MSO_ANCHOR.MIDDLE)
    box(s, Inches(8.1), Inches(1.9), W - MX - Inches(8.1), Inches(4.2), fill=WHITE, line=LINE)
    text(s, Inches(8.4), Inches(2.1), W - MX - Inches(8.7), Inches(3.8),
         [("Rezultatele juriului", {"size": 16, "bold": True, "color": INK3}),
          ("[DE COMPLETAT live]", {"size": 26, "bold": True}),
          ("încercări · falsuri prinse · reguli noi", {"size": 18, "color": INK2}),
          ("„Acest fals a devenit regula #__”", {"size": 20, "bold": True, "color": OK})])


def s_benchmark(prs):
    s = base(prs, "Benchmark: 240 de documente prin agenții reali", "Cifre",
             notes="80 de dosare sintetice, generate de noi: jumătate curate, jumătate scanate prost. Zero date "
                   "personale scăpate, zero oameni cinstiți acuzați, toate falsurile din aceste trei tipuri prinse. "
                   "Spunem corect: pe documente sintetice cu layout cunoscut — arată că lanțul e implementat corect, "
                   "nu că prindem orice fals real. Nu spunem „100% acuratețe”.")
    nums = [("0 / 80", "dosare cu date personale scăpate", OK),
            ("0 / 16", "oameni cinstiți acuzați", OK),
            ("48 / 48", "falsuri prinse: net umflat, sold falsificat, firmă inactivă", INK),
            ("4–5 s", "pe dosar (3 documente), pe un laptop", INK)]
    cw = (W - 2 * MX - Inches(0.9)) / 4
    for i, (v, label, c) in enumerate(nums):
        x = MX + i * (cw + Inches(0.3))
        box(s, x, Inches(1.95), cw, Inches(2.4), fill=WHITE, line=LINE)
        big_number(s, x + Inches(0.25), Inches(2.15), cw - Inches(0.4), v, label, color=c)
    bullets(s, MX, Inches(4.7), W - 2 * MX, Inches(1.8), [
        "Extragere corectă ≥ 92% pe cel mai slab câmp (≥ 95% pe documente curate)",
        "Cost LLM ~0,0006 $ pe dosar · verificarea netului pe date criptate ~0,3 s",
        "Nou: poze de telefon (perspectivă, ±3°, umbră) + CNP greșit, MRZ ≠ CNP, prompt injection — "
        "[DE COMPLETAT cu cifrele rulării]"], size=18, color=INK2)


def s_limits(prs):
    s = base(prs, "Ce nu pretindem", "Onest",
             notes="Juriul apreciază limitele spuse înainte să le găsească ei. Toate sunt în notes/stefan.md.")
    bullets(s, MX, Inches(1.9), W - 2 * MX, Inches(4.6), [
        "Documentele din benchmark sunt sintetice, cu layout cunoscut — nu documente reale de la bănci.",
        "Falsul „perfect” (brut umflat, taxe refăcute corect): semnalul există (netul nu apare în extras), "
        "dar e sub pragul de review.",
        "OCR-ul greșește cifre: înainte de o acuzație, recitim suma de pe imagine; "
        "dacă iese, „de verificat”, nu „fals”.",
        "Criptarea homomorfică nu ascunde nimic de bancă: permite unui terț să calculeze fără să vadă.",
        "ORB nu aprobă credite: detectează fraude și pregătește dosarul. Omul decide."], size=21)


def s_business(prs):
    s = base(prs, "De la demo la producție", "Business case",
             notes="Cifrele de business vin de la mentorii Raiffeisen, cu sursă. Calea de adoptare e fără risc: "
                   "întâi mod umbră, lângă oameni, fără efect asupra deciziilor.")
    card(s, MX, Inches(1.9), Inches(5.8), Inches(2.4), "Impact [DE COMPLETAT cu mentorii]",
         ["• minute de verificare manuală pe dosar → ~1 minut de revizuire",
          "• dosare pe zi · ore și bani economisiți",
          "• pierderi din fraudă de identitate în România"], accent=BRAND, body_size=16)
    card(s, MX, Inches(4.4), Inches(5.8), Inches(2.0), "Reglementare",
         "AI Act: detectarea fraudei financiare e exclusă din „risc ridicat” — avem oricum om în buclă și audit. "
         "GDPR, DORA.", accent=INK, body_size=17)
    x2 = MX + Inches(6.2)
    w2 = W - MX - x2
    text(s, x2, Inches(1.9), w2, Inches(0.5), "Calea de adoptare", size=20, bold=True)
    for i, (t, b) in enumerate([("Mod umbră", "rulează lângă oameni, fără efect"),
                                ("Asistent", "pregătește dosarul, omul decide"),
                                ("Triaj automat", "doar pentru dosarele cu risc mic")]):
        y = Inches(2.5) + i * Inches(0.75)
        box(s, x2, y, w2, Inches(0.62), fill=WHITE, line=LINE)
        text(s, x2 + Inches(0.2), y, w2 - Inches(0.4), Inches(0.62), [(f"{i + 1}. {t} — {b}", {"after": 0})],
             size=17, anchor=MSO_ANCHOR.MIDDLE)
    card(s, x2, Inches(4.75), w2, Inches(1.7), "De ce nu un furnizor existent?",
         "Resistant AI, Inscribe, Onfido sunt în cloud — datele pleacă din bancă. ORB rulează on-prem și leagă "
         "fraudele între dosare.", accent=OK, body_size=14, title_size=18)


def s_team(prs):
    s = base(prs, "Echipa", "Cine ce a făcut",
             notes="Fiecare poate explica tehnic partea lui în 30 de secunde.")
    rows = [["Cine", "Partea"],
            ["Samuel", "orchestrator, Compliance, Arena + self-learning, ZK"],
            ["Narcis", "Paznic (OCR, date personale, seif), Detector, deploy"],
            ["Stefan", "Funcționar, Verificator, criptare homomorfică, ANAF"],
            ["Erin", "frontend (dosar live, Dashboard, Arena), Detectiv + graful de rețele"],
            ["Lorena", "date de test (CNP/IBAN/CUI/MRZ, documentele Mariei), benchmark, pagina ZK, prezentarea"]]
    table(s, MX, Inches(1.9), W - 2 * MX, rows, [0.18, 0.82], size=19, row_h=Inches(0.62))


def s_end(prs):
    s = base(prs, None, notes="Închidere: Maria nu va afla niciodată de la un executor că are o datorie de 45.000 lei. "
                               "Viziune: Radarul care găsește următorul proces de automatizat. Întrebări?", dark=True)
    text(s, MX, Inches(1.8), W - 2 * MX, Inches(2.2),
         "Maria nu va afla niciodată de la un executor că are o datorie de 45.000 lei.", size=40, bold=True,
         color=WHITE)
    text(s, MX, Inches(4.2), W - 2 * MX, Inches(1.2),
         "Următorul pas: Radarul care găsește următorul proces de automatizat.", size=24, color=BRAND)
    text(s, MX, Inches(5.6), W - 2 * MX, Inches(0.8), "Mulțumim! Întrebări?", size=30, bold=True, color=WHITE)


SLIDES = [s_title, s_maria, s_solution, s_agents, s_case, s_privacy, s_zk, s_arena, s_benchmark, s_limits,
          s_business, s_team, s_end]


# ── pitch-ul scurt în engleză (notes/plan.md: 6 slide-uri; textul și cifrele din notes/pitch.md) ──────────────

def en_problem(prs):
    s = base(prs, "Credit back-office work is repetitive", "ORB · The problem",
             notes="Today an analyst reads the documents, recalculates the taxes, looks up the employer at ANAF and "
                   "compares the payslip with the bank statement. Our informal test plus a step-by-step estimate: "
                   "about "
                   "20 minutes per application — and checking for rings takes 5–10 more, or isn't done at all. "
                   "And identity fraud slips through, because rings can't be seen one application at a time.")
    text(s, MX, Inches(1.8), Inches(7), Inches(0.6), "For every loan application, an analyst:", size=22, color=INK2)
    steps = ["Reads the ID card, payslip and bank statement",
             "Recalculates the taxes on the payslip",
             "Looks up the employer at ANAF (the tax authority)",
             "Checks that the net pay arrives in the statement",
             "Searches other files for the same IBAN or phone"]
    for i, t in enumerate(steps):
        y = Inches(2.45) + i * Inches(0.78)
        box(s, MX, y, Inches(0.55), Inches(0.55), fill=BRAND, shape=MSO_SHAPE.OVAL)
        text(s, MX, y, Inches(0.55), Inches(0.55), str(i + 1), size=18, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, MX + Inches(0.8), y, Inches(7.0), Inches(0.55), t, size=18, anchor=MSO_ANCHOR.MIDDLE)
    box(s, Inches(8.7), Inches(1.8), W - MX - Inches(8.7), Inches(4.6), fill=WHITE, line=LINE)
    text(s, Inches(9.0), Inches(2.0), W - MX - Inches(9.3), Inches(4.3),
         [("Manual check, estimated", {"size": 16, "bold": True, "color": INK3}),
          ("~20 min", {"size": 60, "bold": True, "after": 0}),
          ("per application: informal team test + step-by-step estimate (14–21 min)",
           {"size": 15, "color": INK2, "after": 18}),
          ("And fraud rings stay invisible: one application at a time, the same IBAN on 6 identities "
           "looks normal.", {"size": 17, "color": INK2})])


def en_solution(prs):
    s = base(prs, "ORB does the repetitive check in ~20 seconds", "The solution",
             notes="ORB does the repetitive verification in about 20 seconds instead of about 20 minutes, "
                   "without customer "
                   "data leaving the bank. Six specialised agents read and check; every finding is shown on the "
                   "document; the risk level comes from the evidence, not from the AI; the analyst decides and "
                   "signs. Catching fraud is the bonus: on Maria's application ORB finds the net pay that doesn't "
                   "follow from the gross, the inactive employer, the balance that doesn't add up and the same IBAN "
                   "in 6 applications.")
    agents = ["Guard", "Clerk", "Verifier", "Detector", "Detective", "Compliance"]
    gap = Inches(0.18)
    bw = (W - 2 * MX - gap * 5) / 6
    for i, a in enumerate(agents):
        x = MX + i * (bw + gap)
        box(s, x, Inches(1.85), bw, Inches(0.6), fill=INK)
        text(s, x, Inches(1.85), bw, Inches(0.6), a, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    cw = (W - 2 * MX - Inches(0.6)) / 3
    cards = [("Evidence on the document", "the analyst checks the evidence in a minute instead of redoing "
                                               "the whole file", INK),
             ("The AI doesn't decide", "the risk level is computed from the evidence, each one explained; "
                                       "approving a flagged file requires a written reason", OK),
             ("Fraud is the bonus", "Maria: net 8,307 vs 8,500 lei, inactive employer, balance that doesn't add "
                                    "up, same IBAN in 6 applications", DANGER)]
    for i, (t, b, a) in enumerate(cards):
        card(s, MX + i * (cw + Inches(0.3)), Inches(2.75), cw, Inches(2.3), t, b, accent=a, body_size=17,
             title_size=18)
    text(s, MX, Inches(5.4), W - 2 * MX, Inches(1.0),
         [("Also built (in the repo): homomorphic encryption for third-party checks · zero-knowledge income proof · "
           "Arena, where anyone can try to fool the agents · self-learning rules approved by a human",
           {"size": 15, "color": INK3})])


def en_numbers(prs):
    s = base(prs, "The numbers, each with its source", "Numbers",
             notes="Everything ORB does is measured; the manual time is an estimate and we say so. "
                   "We say it openly: synthetic data and "
                   "forgeries made by us — it's a correctness test of the pipeline, not accuracy on real fraud. "
                   "On bad phone photos extraction drops a lot, so ORB sends the file to a human instead of guessing.")
    rows = [["What", "Value", "Source"],
            ["Manual time per application", "~20 min (estimate: 14–21)", "informal team test + steps"],
            ["ORB time per application", "~13–20 s cloud · 8–9 s local", "Maria's file; benchmark"],
            ["AI cost per application", "~$0.005", "LLM usage, Maria's file"],
            ["Real AI (Bedrock), 60 files", "0 to the cloud · 0 invented values", "benchmark --n 60 --llm bedrock"],
            ["Honest people accused (60 files)", "0/12 (3/12 sent to a human)", "idem"],
            ["Forgeries caught (60 files)", "36/36 · fields correct 56–60/60", "idem"],
            ["Phone photos, real AI, 20 files", "0 to cloud · 0/4 accused · 10/12 caught", "--phone --llm bedrock"]]
    table(s, MX, Inches(1.8), W - 2 * MX, rows, [0.33, 0.37, 0.30], size=14, row_h=Inches(0.52))
    box(s, MX, Inches(6.15), W - 2 * MX, Inches(0.55), fill=WARN_BG, line=WARN)
    text(s, MX + Inches(0.2), Inches(6.15), W - 2 * MX - Inches(0.4), Inches(0.55),
         "Synthetic documents and forgeries made by us: a correctness test of the pipeline, "
         "not accuracy on real fraud.",
         size=15, bold=True, color=WARN, anchor=MSO_ANCHOR.MIDDLE)


def en_privacy(prs):
    s = base(prs, "Customer data stays in the bank", "Privacy",
             notes="The AI sees only labels like [PERSON_1] and [CNP_1]. A blurry photo or the ID card never goes to "
                   "the cloud. Or everything runs locally, on one laptop, with no internet. Nothing is absolute: "
                   "these are layers. On the benchmark, zero personal values reached the cloud; the full guarantee "
                   "is local mode.")
    items = [("The AI sees only labels", "names, CNP, IBAN, MRZ are replaced locally by [PERSON_1], [CNP_1] "
                                         "before any AI call", INK),
             ("A gate on the way out", "any payload containing a value from the local vault is blocked; "
                                       "ID card images and blurry photos never leave", INK),
             ("100% local mode", "one switch: the AI runs on the bank's own machine (Ollama), zero external "
                                 "calls — 8–9 s per application", OK),
             ("Every call audited", "each AI call is logged with what was sent; on the benchmark: 0 personal "
                                    "values sent to the cloud", INK)]
    cw = (W - 2 * MX - Inches(0.3)) / 2
    for i, (t, b, a) in enumerate(items):
        card(s, MX + (i % 2) * (cw + Inches(0.3)), Inches(1.9) + (i // 2) * Inches(2.2), cw, Inches(1.95), t, b,
             accent=a, body_size=18)


def en_adoption(prs):
    s = base(prs, "A pilot in shadow mode", "Adoption",
             notes="We measured, we didn't estimate. We propose a pilot in shadow mode: ORB runs next to the analyst "
                   "and decides nothing. Then it pre-fills the checks, then clean files go through fast. The AI "
                   "never makes the decision. Openly: model validation, integration with the credit system and the "
                   "document source are the pilot's steps.")
    steps = [("1", "Shadow mode", "runs next to the analyst, decides nothing; we compare with their results"),
             ("2", "Assistant", "pre-fills the checks and shows the evidence; the analyst decides"),
             ("3", "Triage", "clean, low-risk files go through fast; the rest get full attention")]
    cw = (W - 2 * MX - Inches(0.6)) / 3
    for i, (n, t, b) in enumerate(steps):
        x = MX + i * (cw + Inches(0.3))
        box(s, x, Inches(1.9), cw, Inches(2.5), fill=WHITE, line=LINE)
        box(s, x, Inches(1.9), cw, Inches(0.09), fill=BRAND)
        text(s, x + Inches(0.25), Inches(2.1), cw - Inches(0.5), Inches(2.2),
             [(n, {"size": 40, "bold": True, "after": 0}), (t, {"size": 22, "bold": True}),
              (b, {"size": 16, "color": INK2})])
    card(s, MX, Inches(4.7), (W - 2 * MX - Inches(0.3)) / 2, Inches(1.7), "Human in the loop",
         "AI Act: the AI never decides, every score is explained, every call is audited. GDPR, DORA.",
         accent=OK, body_size=16)
    card(s, MX + (W - 2 * MX + Inches(0.3)) / 2, Inches(4.7), (W - 2 * MX - Inches(0.3)) / 2, Inches(1.7),
         "Pilot steps (said openly)", "model validation (MRM), integration with the credit system, the document "
                                      "source — and labelled real files to measure real fraud",
         accent=INK, body_size=16)


def en_team(prs):
    s = base(prs, "The team", "Who did what",
             notes="Each of us can explain our part technically in 30 seconds.")
    rows = [["Who", "Part"],
            ["Samuel", "orchestrator, Compliance (risk from evidence), Arena, self-learning, ZK"],
            ["Narcis", "Guard (OCR, personal data, vault), Detector, the bank's local machine"],
            ["Stefan", "Clerk, Verifier (taxes, statement, ANAF), homomorphic encryption"],
            ["Erin", "frontend (live case, Dashboard, Arena), Detective + fraud ring graph"],
            ["Lorena", "test data (valid CNP/IBAN/CUI/MRZ), Maria's documents, benchmark, slides"]]
    table(s, MX, Inches(1.9), W - 2 * MX, rows, [0.16, 0.84], size=16, row_h=Inches(0.62))
    text(s, MX, Inches(6.0), W - 2 * MX, Inches(0.5),
         "Live: orb-raiffeisen.vercel.app", size=18, bold=True, color=INK2)


SLIDES_EN = [en_problem, en_solution, en_numbers, en_privacy, en_adoption, en_team]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lang", choices=["ro", "en"], default="ro")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()
    en = args.lang == "en"
    args.out = args.out or str(ROOT / ("prezentare_ORB_EN.pptx" if en else "prezentare_ORB.pptx"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for make in SLIDES_EN if en else SLIDES:
        make(prs)
    prs.save(args.out)
    print(f"{len(prs.slides)} slide-uri scrise în {args.out}")


if __name__ == "__main__":
    main()
